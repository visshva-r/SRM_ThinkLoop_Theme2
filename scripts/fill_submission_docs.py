"""Fill PPT + AI disclosure for final submission."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent.parent.parent
PROJ = Path(__file__).resolve().parent.parent
SUB = PROJ / "submission"

FULL_NAME = "Visshva R"
EMAIL = "vr6123@srmist.edu.in"
TEAM_FORMAL = "SRM_ThinkLoop_Theme2"
COLLEGE = "SRM Institute of Science and Technology"
GITHUB = "https://github.com/visshva-r/SRM_ThinkLoop_Theme2"
LIVE = "https://srm-thinkloop-theme2.onrender.com"
DEMO = "https://github.com/visshva-r/SRM_ThinkLoop_Theme2/blob/main/submission/Demo_Video.mp4"
TODAY = "30 September 2026"


def set_shape_text(shape, text: str, size: int = 14) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = "Calibri"


def fill_body(shape, title: str, body: str, title_size: int = 20, body_size: int = 12) -> None:
    tf = shape.text_frame
    tf.clear()
    p0 = tf.paragraphs[0]
    r0 = p0.add_run()
    r0.text = title
    r0.font.size = Pt(title_size)
    r0.font.bold = True
    r0.font.name = "Calibri"
    for line in body.split("\n"):
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(body_size)
        r.font.name = "Calibri"
        p.space_before = Pt(3)


def fill_ppt() -> Path:
    prs = Presentation(str(ROOT / "CollegeName_TeamName_Submission.pptx"))
    slide1 = prs.slides[0]
    for shape in slide1.shapes:
        if hasattr(shape, "text") and "Theme ID" in (shape.text or ""):
            set_shape_text(
                shape,
                (
                    f"Theme ID - 2 (Guided Troubleshooting)\n"
                    f"Team Name - {TEAM_FORMAL}\n"
                    f"College Name - {COLLEGE}\n"
                    f"Member Name & Email 1- {FULL_NAME} ({EMAIL})\n"
                    f"Member Name & Email 2- N/A (solo)\n"
                    f"Member Name & Email 3- N/A\n"
                    f"Member Name & Email 4- N/A\n"
                    f"Submission Github link - {GITHUB}"
                ),
                size=12,
            )

    slides_content = {
        2: (
            "Theme",
            "Theme 2: Smart Guided Troubleshooting Engine\n"
            "Problem: Convert a natural-language Galaxy issue + SIIS support text into a "
            "schema-valid troubleshooting plan with masked Settings deeplinks.\n"
            "Constraints: no URL leaks, grounded steps from SIIS only, cache for repeats/"
            "paraphrases, generalize to unseen SIIS payloads, live API for judges.",
        ),
        3: (
            "Existing Solutions & Gaps",
            "Existing: generic LLM chatbots, static FAQ pages, SIIS text dumps without structure.\n"
            "Gaps we close:\n"
            "- No schema enforcement → invalid goals/titles/descriptions fail automated gates\n"
            "- Raw URLs / open links → G5 disqualification\n"
            "- No cache → repeat/paraphrase latency fails A3\n"
            "- Invented steps → not grounded in SIIS\n"
            "- No deeplink catalog matching → auto actions missing actionable deeplinks",
        ),
        4: (
            "Our Solution & Architecture",
            "Live FastAPI: GET /health  ·  POST /v1/troubleshoot\n"
            "Pipeline: Normalize → Cache → Gemini Stage1 plan → Grounding → "
            "Hybrid retrieval (BM25 + embeddings locally; BM25-only LIGHTWEIGHT on Render) → "
            "Gemini Stage2 deeplink pick → Repair/validate → JSON\n"
            "Fallback: extractive SIIS Step parser when LLM fails/timeouts\n"
            "Warm cache from results.jsonl (exact + paraphrase keys)\n"
            f"Deploy: Docker on Render free tier  ·  {LIVE}",
        ),
        5: (
            "Demo & Product Walkthrough",
            "Demo video in repo: submission/Demo_Video.mp4\n"
            f"{DEMO}\n"
            "Walkthrough shown:\n"
            '1) /health → {"status":"ok"}\n'
            "2) Kit blank-display query → structured contexts + bixby:// deeplinks\n"
            "3) Exact repeat → meta.cache_hit=true\n"
            "4) Paraphrase → cache hit from warmed variations\n"
            "5) Unseen Wi-Fi SIIS → gemini-2.5-flash, grounded steps, no URL leaks\n"
            f"Swagger UI: {LIVE}/docs",
        ),
        6: (
            "Tools and Tech Stack",
            "Language / API: Python 3.11, FastAPI, Uvicorn, Pydantic v2\n"
            "LLM: Google Gemini 2.5 Flash (google-genai) — Stage1 plan + Stage2 pick\n"
            "Retrieval: rank-bm25, fastembed ONNX (local); LIGHTWEIGHT BM25 on Render\n"
            "Cache: in-memory exact + paraphrase keys; semantic embeddings locally\n"
            "Data: deeplinks.json (578), siis_responses.json (20), results.jsonl\n"
            "Deploy: Docker, Render free, render.yaml + Dockerfile\n"
            "Eval: scripts/eval_local.py (G2–G5, A1–A5), pytest",
        ),
        7: (
            "Impact & Use Case",
            "Use case: Samsung support / Bixby-style guided fix flows for Galaxy devices.\n"
            "User describes a blank screen, Wi-Fi failure, cracked display, etc.\n"
            "System returns ordered Goals → Actions → StepGroups with masked deeplinks "
            "so an assistant can open the right Settings screen safely.\n"
            "Impact: faster first-response troubleshooting, fewer invented steps, "
            "judge-safe outputs (schema + zero URL leaks), reproducible Docker deploy.",
        ),
        8: (
            "Innovation, Results & Limitations",
            "Innovation: hedged JSON shape + repair layer; SIIS grounding; dummy deeplink "
            "policy; warm cache from offline results; Render LIGHTWEIGHT_MODE under 512MB.\n"
            "Results: G2–G5 PASS; A1 15/15; A2 15/15; A5 5/5; A4 8/10; "
            "live black-box checks 6/6; cache_hit 100% on kit/paraphrase.\n"
            "Limitations: Render free network latency (~700ms) vs 300ms A3 target; "
            "cold Gemini path slow on free tier; dense retrieval disabled on free RAM.",
        ),
        9: (
            "What’s Next",
            "- Re-enable hybrid dense retrieval on a larger host for better deeplink precision\n"
            "- Tighten cold-path p95 with response streaming / hedged timeouts\n"
            "- Expand paraphrase embedding cache beyond exact keys on production memory\n"
            "- Add more SIIS domains beyond the 20-kit scenarios\n"
            "- Package as a PRISM worklet with Samsung mentor feedback",
        ),
        10: (
            "Differentiation (Brownie Points)",
            "- Gemini + extractive fallback (works even if LLM is down)\n"
            "- Programmatic repair for goal/title/description/category rules\n"
            "- Zero URL leak scrubbing across the full JSON tree\n"
            "- Cache warm-load so kit + variations hit immediately after deploy\n"
            "- Free-tier deploy that still passes live health/schema/generalization checks\n"
            "- Local scorer reproducing G2–G5 and A1–A5 for iteration speed",
        ),
        11: (
            "Checklist — Updated on Public GitHub",
            "Working prototype code — public GitHub repo ………… YES\n"
            "README with reproducible setup + Docker …………… YES\n"
            "Demo video ≤5 min (in repo) ………………………… YES\n"
            "  → submission/Demo_Video.mp4\n"
            "Presentation file (PPT) ……………………………… YES\n"
            "  → submission/SRM_ThinkLoop_Submission.pptx\n"
            "AI Disclosure ………………………………………… YES\n"
            "  → submission/LangAI3.0_AI_Disclosure.docx\n"
            "Release tag PRISM_GENAI_HACKATHON_Y2026 …… YES (retagged on final commit)\n"
            f"Live API ……………………………………………… YES — {LIVE}",
        ),
    }

    for idx, (title, body) in slides_content.items():
        slide = prs.slides[idx - 1]
        candidates = [s for s in slide.shapes if hasattr(s, "text_frame")]
        target = None
        for s in candidates:
            text = s.text or ""
            if (
                title.split()[0].lower() in text.lower()
                or "Checklist" in text
                or "Working prototype" in text
            ):
                target = s
                break
        if target is None and candidates:
            target = max(candidates, key=lambda s: s.width * s.height)
        if target is not None:
            fill_body(target, title, body)

    out = SUB / "SRM_ThinkLoop_Submission.pptx"
    prs.save(str(out))
    return out


def fill_disclosure() -> Path:
    doc = Document()
    doc.add_heading("AI Usage DISCLOSURE FORM", level=1)
    doc.add_heading("1. Team Details", level=2)
    doc.add_paragraph(f"Team Name: {TEAM_FORMAL}")
    doc.add_paragraph("Project / Product Name: Smart Guided Troubleshooting Engine")
    doc.add_paragraph(f"Organization / Institution: {COLLEGE}")
    doc.add_paragraph(f"Submission Date: {TODAY}")

    doc.add_heading("2. AI Usage Declaration", level=2)
    doc.add_paragraph(
        "Did your team use any Artificial Intelligence (AI) in developing this project?  Yes"
    )

    doc.add_heading("3. Purpose of AI Usage (Brief Details)", level=2)
    doc.add_paragraph(
        "Idea generation / brainstorming: Yes — architecture options and Theme 2 scoring "
        "strategy discussed with Cursor AI assistants."
    )
    doc.add_paragraph(
        "Code generation or assistance: Yes — FastAPI pipeline, retrieval, cache, repair, "
        "eval scripts, Docker/Render deploy assisted in Cursor (Composer / Auto) with human "
        "review and testing."
    )
    doc.add_paragraph(
        "UI / UX design: No dedicated product UI; Swagger UI is FastAPI default. "
        "No AI-designed frontend."
    )
    doc.add_paragraph(
        "Content creation: Yes — README, architecture notes, demo script, and presentation "
        "content drafted with AI assistance and edited by the team."
    )
    doc.add_paragraph(
        "Data analysis: Limited — local eval metrics and gate scoring scripts; kit data is "
        "organizer-provided."
    )
    doc.add_paragraph(
        "Testing / debugging: Yes — diagnosing Render OOM, Gemini response shape issues, "
        "and black-box API checks with AI assistance."
    )
    doc.add_paragraph(
        "Other: Runtime LLM in the product itself — Google Gemini 2.5 Flash for Stage-1 "
        "plan extraction and Stage-2 deeplink selection (SIIS grounding + extractive "
        "fallback applied; steps are not freely invented)."
    )

    doc.add_heading("4. Feature Origin Classification", level=2)
    features = [
        (
            "FastAPI /health and /v1/troubleshoot service",
            "Both",
            "Tools: Cursor Agent (Composer/Auto). Theme: implement Theme 2 REST API matching "
            "student kit schema. Output: app/main.py and pipeline. Modified: hedged response "
            "shape, Render PORT handling, human-reviewed endpoints.",
        ),
        (
            "Gemini Stage-1 plan + Stage-2 deeplink picker",
            "Both",
            "Tools: Cursor + Google Gemini API (gemini-2.5-flash) at runtime. Prompts constrain "
            "JSON nesting and SIIS-only steps. Output: app/llm.py. Modified: "
            "_normalize_stage1_plan for Gemini layout mistakes; fallback on failure.",
        ),
        (
            "Extractive SIIS fallback planner",
            "Both",
            "Tools: Cursor. Output: app/fallback.py Step/section parser. Modified: tuned for "
            "kit markdown; used when LLM absent/fails so gates still pass.",
        ),
        (
            "Hybrid retrieval + LIGHTWEIGHT_MODE cache",
            "Both",
            "Tools: Cursor. Output: BM25 + fastembed locally; BM25-only + exact/paraphrase "
            "warm cache on Render 512MB after OOM. Modified: LIGHTWEIGHT_MODE for free tier.",
        ),
        (
            "Schema repair, URL scrub, local scorer",
            "Both",
            "Tools: Cursor. Output: app/repair.py, scripts/eval_local.py. Modified: aligned "
            "with FAQ gates G2–G5 and blocks A1–A5; verified with pytest and live checks.",
        ),
    ]
    for i, (name, origin, desc) in enumerate(features, 1):
        doc.add_paragraph(f"{i}. Feature Name: {name}")
        doc.add_paragraph(f"   Origin: {origin}")
        doc.add_paragraph(f"   Description: {desc}")

    doc.add_heading("5. Ethical & Compliance Confirmation", level=2)
    doc.add_paragraph("AI usage complies with guidelines and policies: Yes")
    doc.add_paragraph("No proprietary or copyrighted data misused: I Agree")
    doc.add_paragraph(
        "Organizer student-kit files (schema, SIIS, deeplinks) used only as permitted for "
        "the hackathon. No real customer PII. No leaked http(s) URLs in API outputs."
    )

    doc.add_heading("6. Declaration & Sign-Off", level=2)
    doc.add_paragraph(f"Name of Team Representative: {FULL_NAME}")
    doc.add_paragraph("Role: Team Lead / Sole Member")
    doc.add_paragraph(f"Signature: {FULL_NAME} (digital)")
    doc.add_paragraph(f"Date: {TODAY}")
    doc.add_paragraph(f"Contact Email: {EMAIL}")

    out = SUB / "LangAI3.0_AI_Disclosure.docx"
    doc.save(str(out))
    return out


def main() -> None:
    SUB.mkdir(parents=True, exist_ok=True)
    ppt = fill_ppt()
    disc = fill_disclosure()
    print(ppt)
    print(disc)


if __name__ == "__main__":
    main()
