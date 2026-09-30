"""Create architecture diagram PNG and embed it in the submission PPT slide 4."""
from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

PROJ = Path(__file__).resolve().parent.parent
SUB = PROJ / "submission"
PNG = SUB / "architecture_diagram.png"
PPTX = SUB / "SRM_ThinkLoop_Submission.pptx"


def _font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    return ImageFont.load_default()


def draw_box(draw, xy, text, fill, outline, font, text_fill=(20, 30, 45)):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=2)
    bbox = draw.textbbox((0, 0), text, font=font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    tx = x1 + (x2 - x1 - tw) / 2
    ty = y1 + (y2 - y1 - th) / 2
    draw.text((tx, ty), text, fill=text_fill, font=font)


def draw_arrow(draw, x1, y1, x2, y2, color=(70, 90, 120)):
    draw.line((x1, y1, x2, y2), fill=color, width=3)
    # arrow head
    if x2 >= x1:
        draw.polygon([(x2, y2), (x2 - 10, y2 - 6), (x2 - 10, y2 + 6)], fill=color)
    else:
        draw.polygon([(x2, y2), (x2 + 10, y2 - 6), (x2 + 10, y2 + 6)], fill=color)


def make_diagram() -> Path:
    w, h = 1600, 900
    img = Image.new("RGB", (w, h), (248, 250, 252))
    draw = ImageDraw.Draw(img)
    title_f = _font(36, bold=True)
    box_f = _font(22, bold=True)
    small_f = _font(18)
    tiny_f = _font(16)

    draw.text((40, 28), "Smart Guided Troubleshooting — Architecture", fill=(30, 45, 70), font=title_f)
    draw.text(
        (40, 78),
        "Query + SIIS  →  Cache / Gemini plan  →  Ground  →  Retrieve  →  Pick deeplink  →  Repair  →  JSON",
        fill=(80, 95, 120),
        font=small_f,
    )

    # Row 1 main pipeline
    boxes = [
        (40, 160, 220, 250, "Client\nPOST /v1", (220, 235, 250), (70, 110, 180)),
        (260, 160, 440, 250, "Normalize\n+ Cache", (220, 245, 230), (40, 140, 100)),
        (480, 160, 700, 250, "Gemini Stage1\nPlan + variations", (255, 240, 220), (190, 120, 40)),
        (740, 160, 920, 250, "Grounding\nSIIS filter", (240, 230, 255), (120, 80, 180)),
        (960, 160, 1180, 250, "Retrieval\nBM25 (+embed)", (220, 245, 255), (40, 120, 170)),
        (1220, 160, 1440, 250, "Gemini Stage2\nDeeplink pick", (255, 240, 220), (190, 120, 40)),
        (1480 - 40, 160, 1560, 250, "Repair\n+ JSON", (255, 230, 230), (170, 60, 70)),
    ]
    # Fix last box width - use cleaner layout with 7 boxes
    boxes = [
        (30, 170, 200, 270, "Client API\n/health\n/troubleshoot", (220, 235, 250), (70, 110, 180)),
        (230, 170, 420, 270, "Normalize\n+ Semantic\nCache", (220, 245, 230), (40, 140, 100)),
        (450, 170, 660, 270, "Gemini S1\nPlan extract\n+ variations", (255, 240, 220), (190, 120, 40)),
        (690, 170, 870, 270, "Grounding\nSIIS-only\nsteps", (240, 230, 255), (120, 80, 180)),
        (900, 170, 1100, 270, "Hybrid\nRetrieval\nBM25+embed*", (220, 245, 255), (40, 120, 170)),
        (1130, 170, 1340, 270, "Gemini S2\nPick catalog\ndeeplink", (255, 240, 220), (190, 120, 40)),
        (1370, 170, 1570, 270, "Repair &\nValidate\n→ JSON", (255, 230, 230), (170, 60, 70)),
    ]

    for x1, y1, x2, y2, text, fill, outline in boxes:
        draw_box(draw, (x1, y1, x2, y2), text, fill, outline, box_f)

    for i in range(len(boxes) - 1):
        x1 = boxes[i][2]
        x2 = boxes[i + 1][0]
        mid_y = (boxes[i][1] + boxes[i][3]) // 2
        draw_arrow(draw, x1 + 4, mid_y, x2 - 4, mid_y)

    # Cache hit bypass
    draw.line((325, 270, 325, 340), fill=(40, 140, 100), width=3)
    draw.line((325, 340, 1470, 340), fill=(40, 140, 100), width=3)
    draw.line((1470, 340, 1470, 270), fill=(40, 140, 100), width=3)
    draw.polygon([(1470, 270), (1464, 282), (1476, 282)], fill=(40, 140, 100))
    draw.text((700, 350), "cache hit → return cached plan (fast path)", fill=(40, 140, 100), font=small_f)

    # Fallback box
    draw_box(
        draw,
        (450, 420, 870, 520),
        "Extractive Fallback\nSIIS Step parser (no LLM)",
        (255, 248, 220),
        (160, 130, 40),
        box_f,
    )
    draw.line((555, 270, 555, 420), fill=(160, 130, 40), width=3)
    draw.polygon([(555, 420), (549, 408), (561, 408)], fill=(160, 130, 40))
    draw.text((570, 300), "LLM fail / timeout", fill=(160, 130, 40), font=tiny_f)
    draw.line((870, 470, 1000, 470), fill=(160, 130, 40), width=3)
    draw.line((1000, 470, 1000, 270), fill=(160, 130, 40), width=3)

    # Data stores
    draw_box(
        draw,
        (100, 580, 500, 700),
        "Data\ndeeplinks.json (578)\nsiis_responses.json (20)\nresults.jsonl warm-load",
        (245, 245, 250),
        (90, 100, 130),
        small_f,
    )
    draw_box(
        draw,
        (560, 580, 1000, 700),
        "Deploy\nDocker · Render free\nLIGHTWEIGHT_MODE=true\n(BM25-only, no embed RAM)",
        (245, 245, 250),
        (90, 100, 130),
        small_f,
    )
    draw_box(
        draw,
        (1060, 580, 1500, 700),
        "Outputs\nGoals → Actions → Steps\nbixby:// masked deeplinks\nmeta: cache_hit, model, latency",
        (245, 245, 250),
        (90, 100, 130),
        small_f,
    )

    draw.text(
        (40, 740),
        "* Local: BM25 + ONNX embeddings. Render free (512MB): BM25-only LIGHTWEIGHT_MODE.",
        fill=(90, 100, 120),
        font=tiny_f,
    )
    draw.text(
        (40, 780),
        "Live: https://srm-thinkloop-theme2.onrender.com   ·   Team: SRM_ThinkLoop_Theme2",
        fill=(70, 90, 130),
        font=small_f,
    )
    draw.text(
        (40, 840),
        "Samsung PRISM Gen AI Hackathon 3.0 — Theme 2",
        fill=(120, 130, 150),
        font=tiny_f,
    )

    SUB.mkdir(parents=True, exist_ok=True)
    img.save(PNG, "PNG")
    return PNG


def embed_in_ppt() -> None:
    prs = Presentation(str(PPTX))
    slide = prs.slides[3]  # slide 4 (0-indexed)

    # Shrink / replace main text body
    candidates = [s for s in slide.shapes if hasattr(s, "text_frame")]
    target = max(candidates, key=lambda s: s.width * s.height) if candidates else None
    if target is not None:
        tf = target.text_frame
        tf.clear()
        p0 = tf.paragraphs[0]
        r0 = p0.add_run()
        r0.text = "Our Solution & Architecture Diagram"
        r0.font.size = Pt(20)
        r0.font.bold = True
        r0.font.name = "Calibri"
        p1 = tf.add_paragraph()
        r1 = p1.add_run()
        r1.text = (
            "FastAPI pipeline with Gemini planning, SIIS grounding, hybrid retrieval, "
            "repair/validate, extractive fallback, and warm cache. Live on Render (Docker)."
        )
        r1.font.size = Pt(11)
        r1.font.name = "Calibri"
        # Move text box to top strip if possible
        target.top = Inches(0.35)
        target.left = Inches(0.4)
        target.width = Inches(9.2)
        target.height = Inches(0.95)

    # Add diagram image
    slide.shapes.add_picture(str(PNG), Inches(0.25), Inches(1.35), width=Inches(9.5))
    prs.save(str(PPTX))


def main() -> None:
    path = make_diagram()
    print("Wrote", path)
    embed_in_ppt()
    print("Updated", PPTX)


if __name__ == "__main__":
    main()
